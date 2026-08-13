from __future__ import annotations

"""Compare the deployed coefficient lookup against the 50k planning workbook.

Outputs aggregate mapping evidence only; no customer-level data is exported.
"""

import json
from collections import Counter
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook

DESKTOP = Path("C:/Users/PP-2026070302/Desktop")
COEFFICIENT_FILE = DESKTOP / "\u63d0\u989d\u7cfb\u657020260812.xlsx"
PLAN_FILE = DESKTOP / "\u5927\u6a21\u578b\u6837\u672c2-\u63d0\u989d\u65b9\u6848-ABC\u5ba2\u7fa4-20260812.xlsx"
OUTPUT = Path("risk_analysis/strategy_coefficient_mapping.json")


def number(value: object) -> float | None:
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def key(row: tuple[object, ...]) -> tuple[str, str, int, str]:
    return (str(row[0]).strip(), str(row[1]).strip(), int(float(row[2])), str(row[3]).strip())


def main() -> None:
    coefficient_wb = load_workbook(COEFFICIENT_FILE, data_only=True, read_only=True)
    coefficient_ws = coefficient_wb[coefficient_wb.sheetnames[0]]
    coeff: dict[tuple[str, str, int, str], float] = {}
    duplicates: list[str] = []
    for row in coefficient_ws.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue
        lookup_key = key(row)
        if lookup_key in coeff:
            duplicates.append(str(lookup_key))
        coeff[lookup_key] = float(row[4])

    plan_wb = load_workbook(PLAN_FILE, data_only=True, read_only=True)
    plan_ws = plan_wb["\u63d0\u989d\u7ed3\u679c"]
    plan: dict[tuple[str, str, int, str], dict[str, float | bool | None]] = {}
    for row in plan_ws.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue
        lookup_key = key(row)
        plan[lookup_key] = {
            "eligible_abc": bool(row[18]),
            "raw_cap": number(row[16]),
            "plan_a_rate": number(row[19]),
            "plan_b_rate": number(row[22]),
            "plan_a_after_limit": number(row[20]),
        }

    joined = []
    for lookup_key, deployed_rate in coeff.items():
        planned = plan.get(lookup_key)
        joined.append(
            {
                "key": "|".join(map(str, lookup_key)),
                "deployed_rate": deployed_rate,
                **(planned or {}),
            }
        )

    relation = Counter()
    eligible_counts = Counter()
    scale_factors: list[float] = []
    deviations: list[dict[str, object]] = []
    for item in joined:
        if "raw_cap" not in item:
            relation["missing_from_50k_plan"] += 1
            continue
        raw = item["raw_cap"]
        plan_a = item["plan_a_rate"]
        plan_b = item["plan_b_rate"]
        deployed = item["deployed_rate"]
        eligible_counts[str(item["eligible_abc"])] += 1
        if raw is not None and abs(deployed - raw) < 1e-9:
            relation["equals_raw_cap"] += 1
        elif plan_b is not None and abs(deployed - plan_b) < 1e-9:
            relation["equals_plan_b"] += 1
        elif plan_a is not None and abs(deployed - plan_a) < 1e-9:
            relation["equals_plan_a"] += 1
        else:
            relation["other"] += 1
            deviations.append(item)
        if raw not in (None, 0) and bool(item["eligible_abc"]):
            scale_factors.append(float(deployed) / float(raw))

    output = {
        "coefficient_row_count": len(coeff),
        "coefficient_duplicate_keys": duplicates,
        "plan_row_count": len(plan),
        "joined_count": len(joined),
        "mapping_relation": relation,
        "eligible_combination_count": eligible_counts,
        "deployed_to_raw_cap_scale_for_eligible": {
            "min": min(scale_factors) if scale_factors else None,
            "max": max(scale_factors) if scale_factors else None,
            "mean": mean(scale_factors) if scale_factors else None,
            "distribution": Counter(round(item, 6) for item in scale_factors),
        },
        "deviations_sample": deviations[:30],
        "coefficient_distribution": Counter(coeff.values()),
    }
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
