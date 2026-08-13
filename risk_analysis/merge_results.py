"""合并三个结果 Excel → 一个汇总 Excel(保留样式/格式/条件格式/列宽)"""
import os, sys
from copy import copy
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

OUT_DIR = "/data2/jupyter-wenning/strage"
SRC = [
    ("分析", OUT_DIR + "/短账龄客户-模型效果与提额策略分析-20260807_2343.xlsx",
     ["模型效果", "最优分箱", "等频分箱", "提额建议"]),
    ("影响", OUT_DIR + "/提额影响评估-20260807.xlsx",
     ["总体影响", "分客群明细"]),
    ("方案", OUT_DIR + "/提额6000万方案-20260807.xlsx",
     ["方案总览", "达标组合", "分层明细", "客群明细"]),
]
OUT_PATH = OUT_DIR + "/风控提额分析汇总-20260808.xlsx"

wb_out = Workbook()
wb_out.remove(wb_out.active)

def copy_sheet(ws_src, ws_dst):
    # 单元格值与样式
    for row in ws_src.iter_rows():
        for cell in row:
            if cell.value is None and not cell.has_style:
                continue
            nc = ws_dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                nc._style = copy(cell._style)
    # 列宽/行高
    for k, v in ws_src.column_dimensions.items():
        ws_dst.column_dimensions[k].width = v.width
    for k, v in ws_src.row_dimensions.items():
        ws_dst.row_dimensions[k].height = v.height
    # 合并单元格
    for rng in ws_src.merged_cells.ranges:
        ws_dst.merge_cells(str(rng))
    # 条件格式
    for cf in ws_src.conditional_formatting:
        for rule in cf.rules:
            ws_dst.conditional_formatting.add(str(cf.sqref), copy(rule))
    # 冻结/筛选
    if ws_src.freeze_panes:
        ws_dst.freeze_panes = ws_src.freeze_panes
    if ws_src.auto_filter.ref:
        ws_dst.auto_filter.ref = ws_src.auto_filter.ref

for prefix, path, sheets in SRC:
    print(f"加载: {os.path.basename(path)}")
    wb_src = load_workbook(path)
    for s in sheets:
        if s not in wb_src.sheetnames:
            print(f"  跳过(不存在): {s}")
            continue
        name = f"{prefix}_{s}"
        ws_dst = wb_out.create_sheet(name)
        copy_sheet(wb_src[s], ws_dst)
        print(f"  ✓ {name}: {ws_dst.max_row:,} 行 × {ws_dst.max_column} 列")
    wb_src.close()

wb_out.save(OUT_PATH)
print(f"\n✅ 汇总 Excel 已生成: {OUT_PATH}")
print("Sheets:", wb_out.sheetnames)
