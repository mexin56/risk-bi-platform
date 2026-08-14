# -*- coding: utf-8 -*-
"""PDL 提额方案最终版 —— 高280%/中260%/低60% + 风险系数矩阵
思路: 额度使用率优先(使用意愿) × 风险档调节; 低使用率不补量, 中用补量; 压降最大提幅
"""
import warnings; warnings.filterwarnings("ignore")
import pandas as pd, numpy as np, os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

SRC = "/data2/jupyter-wenning/strage/PDL提额系数20260814.xlsx"
DST = "/data2/jupyter-wenning/strage/PDL提额方案-最终版-20260814.xlsx"
TARGET_AMT = 50000.0
CAP_QUOTA = 500000.0
USE_BASE = {"高额度使用率": 2.8, "中额度使用率": 2.6, "低额度使用率": 0.6}
RISK_COEF = [
    (0.02, 1.00, "T1_<2%"), (0.04, 0.95, "T2_2-4%"), (0.06, 0.90, "T3_4-6%"),
    (0.08, 0.85, "T4_6-8%"), (0.10, 0.80, "T5_8-10%"), (0.15, 0.70, "T6_10-15%"),
    (0.20, 0.60, "T7_15-20%"), (0.25, 0.50, "T8_20-25%"), (0.30, 0.40, "T9_25-30%"),
    (np.inf, 0.30, "T10_>=30%"),
]
def rc_of(f):
    for cap, c, _ in RISK_COEF:
        if f < cap:
            return c
    return 0.30
def tier_of(f):
    for cap, _, n in RISK_COEF:
        if f < cap:
            return n
    return "T10_>=30%"

# ---- 1. 读数据 ----
df = pd.read_excel(SRC, sheet_name=0)
cols, seen = [], {}
for c in df.columns:
    s = str(c).strip()
    seen[s] = seen.get(s, 0) + 1
    cols.append(s if seen[s] == 1 else f"{s}_数值")
df.columns = cols
df = df[df["flag_customer"] != "总计"].reset_index(drop=True)
assert (df["结清客户"] - df["提额客户"] - df["未提额客户"]).abs().max() <= 1

N = df["结清客户"].sum()
cur_amt = (df["结清客户"] * df["平均额度"]).sum()
need = TARGET_AMT * N - cur_amt
cur_fpd7 = (df["结清客户"] * df["平均额度"] * df["fpd7_rate$"]).sum() / cur_amt
print(f"PDL结清客户 {N:,} | 当前平均 {cur_amt/N:,.0f} | 大盘fpd7_rate$ {cur_fpd7*100:.2f}% | 需新增 {need/1e8:.2f} 亿")

df["等级可提"] = df["评级"].isin(["A", "B", "C"]) | ((df["评级"] == "D") & (df["fpd7_rate$"] < cur_fpd7))
df["是否可提额"] = df["等级可提"]
df["风险档"] = df["fpd7_rate$"].apply(tier_of)
df["风险系数"] = df["fpd7_rate$"].apply(rc_of)
df["使用率基础"] = df["use_rate_group"].map(USE_BASE)
df["基础提幅"] = np.minimum(df["使用率基础"] * df["风险系数"], 3.0)

# ---- 2. 计算 ----
def compute(k=1.0):
    amp = np.where(df["是否可提额"], df["基础提幅"] * k, 0.0)
    new_per = np.minimum(df["提额客户_平均额度"] * (1 + amp), CAP_QUOTA)
    inc = df["提额客户"] * np.maximum(0, new_per - df["提额客户_平均额度"])
    return amp, new_per, inc

amp_b, new_per_b, inc_b = compute(1.0)
avg_b = (cur_amt + inc_b.sum()) / N
print(f"方案B[满档位不缩放]: 新增 {inc_b.sum()/1e8:.2f} 亿 | 可达 {avg_b:,.0f}")

lo, hi = 0.2, 1.0
for _ in range(80):
    mid = (lo + hi) / 2
    if compute(mid)[2].sum() > need:
        hi = mid
    else:
        lo = mid
k = (lo + hi) / 2
amp_a, new_per_a, inc_a = compute(k)
avg_a = (cur_amt + inc_a.sum()) / N
print(f"方案A[×{k:.3f}]: 新增 {inc_a.sum()/1e8:.2f} 亿 | 平均额度 {avg_a:,.0f} ✅")

df["方案A幅度"] = amp_a
df["方案A后额度"] = new_per_a
df["方案A新增"] = inc_a
df["方案A新增(万)"] = (inc_a / 1e4).round(0)
df["方案B幅度"] = amp_b
df["方案B后额度"] = new_per_b
df["方案B新增"] = inc_b

# ---- 3. 决策依据 ----
def reason(r):
    if not r["是否可提额"]:
        if r["评级"] == "E":
            return "评级=E, E级不予提额"
        return f"评级={r['评级']}, fpd7_rate${r['fpd7_rate$']*100:.1f}%≥大盘{cur_fpd7*100:.1f}%, 不予提额"
    return (f"评级={r['评级']}, fpd7_rate${r['fpd7_rate$']*100:.1f}%({r['风险档']}), {r['use_rate_group']}"
            f"(基础{USE_BASE[r['use_rate_group']]*100:.0f}%)×风险系数{r['风险系数']:.2f}×{k:.3f}→+{r['方案A幅度']*100:.0f}%, "
            f"提额后额度{r['方案A后额度']:,.0f}元")
df["提额决策依据"] = df.apply(reason, axis=1)

# ---- 4. 整体指标 ----
amt_after = df["结清客户"] * df["平均额度"] + df["方案A新增"]
fpd7s_a = (amt_after * df["fpd7_rate$"]).sum() / amt_after.sum()
amt_after_b = df["结清客户"] * df["平均额度"] + df["方案B新增"]
fpd7s_b = (amt_after_b * df["fpd7_rate$"]).sum() / amt_after_b.sum()

# ---- 5. 提幅矩阵 ----
order = [t[2] for t in RISK_COEF]
matrix_rows = []
for u in ["高额度使用率", "中额度使用率", "低额度使用率"]:
    row = {"使用率": u, "基础档%": USE_BASE[u] * 100}
    for tg in order:
        sub = df[(df["use_rate_group"] == u) & (df["风险档"] == tg) & df["是否可提额"]]
        if len(sub):
            row[tg] = round((sub["提额客户"] * sub["方案A幅度"]).sum() / sub["提额客户"].sum() * 100, 0)
        else:
            row[tg] = None
    matrix_rows.append(row)
matrix_df = pd.DataFrame(matrix_rows)

# ---- 6. 思路 sheet ----
idea = pd.DataFrame([
    {"模块": "方案目标", "内容": "整体平均额度从 42,097 提升到 50,000 元, 同时整体 fpd7_rate$ 下降"},
    {"模块": "提额资格", "内容": "评级 A/B/C 全部 + D级且客群fpd7_rate$<大盘(21.25%)可提额; E级一律不提额"},
    {"模块": "提额原则①", "内容": "额度使用率优先(使用意愿): 高使用率基础提幅280% > 中使用率260% > 低使用率60%, 给有使用意愿的客户提额"},
    {"模块": "提额原则②", "内容": "风险档调节: 客群fpd7_rate$分10档, 风险系数×1.00(超低风险)~×0.30(高风险), 风险越高提幅越低"},
    {"模块": "提额原则③", "内容": "低使用率客户使用意愿低, 不承担补量任务(仅贡献2.5%); 补量由中使用率客群承担(中档260%)"},
    {"模块": "提额原则④", "内容": "压降最大提幅: 最终最大提幅约265%(原设计300%封顶压降35pp)"},
    {"模块": "提额公式", "内容": f"提幅 = min(使用率基础档 × 风险系数 × {k:.3f}, 300%); 提额后额度 = min(原额度×(1+提幅), 500,000)"},
    {"模块": "约束条件", "内容": "未提额客户额度保持不变; 结清客户 = 提额客户 + 未提额客户(勾兑校验通过)"},
    {"模块": "达标测算", "内容": f"新增授信 {inc_a.sum()/1e8:.2f} 亿(+{inc_a.sum()/cur_amt*100:.1f}%), 整体平均额度 {avg_a:,.0f}元, fpd7_rate$ {cur_fpd7*100:.2f}%→{fpd7s_a*100:.2f}%"},
])

# ---- 7. 汇总 ----
summary = pd.DataFrame([
    {"模块": "方案目标", "指标": "整体平均额度", "数值": f"{cur_amt/N:,.0f} → {avg_a:,.0f} 元",
     "说明": f"目标 {TARGET_AMT:,.0f} 元, 需新增 {need/1e8:.2f} 亿 ✅ 方案A达标"},
    {"模块": "约束规则", "指标": "提额客群范围", "数值": "A/B/C 全部 + D级低风险",
     "说明": f"D级且fpd7_rate$<大盘{cur_fpd7*100:.1f}%可提额; E级({df.loc[df['评级']=='E','提额客户'].sum():,})及D级高风险({df.loc[(df['评级']=='D')&~df['是否可提额'],'提额客户'].sum():,})不予提额"},
    {"模块": "提幅规则", "指标": "额度使用率优先", "数值": "高280% / 中260% / 低60%",
     "说明": "低使用率不补量(贡献2.5%), 中使用率承担补量"},
    {"模块": "提幅规则", "指标": "风险档系数", "数值": "10档 ×1.00 ~ ×0.30",
     "说明": "T1<2%×1.00 / T3 4-6%×0.90 / T5 8-10%×0.80 / T7 15-20%×0.60 / T9 25-30%×0.40 / T10>=30%×0.30"},
    {"模块": "提幅规则", "指标": "最终提幅", "数值": f"min(使用率基础×风险系数×{k:.3f}, 300%)",
     "说明": f"最大提幅约{np.max(np.where(df['是否可提额'], df['方案A幅度'], 0))*100:.0f}%, 满档位可达{avg_b:,.0f}, 缩放至精确50,000"},
    {"模块": "勾兑关系", "指标": "结清/提额/未提额", "数值": f"{N:,} = {df['提额客户'].sum():,} + {df['未提额客户'].sum():,}",
     "说明": "勾兑校验通过, 未提额客户额度保持不变"},
    {"模块": "方案A(达标5万)", "指标": "整体平均额度", "数值": f"{avg_a:,.0f} 元", "说明": "✅ 精确达标"},
    {"模块": "方案A(达标5万)", "指标": "新增总授信", "数值": f"{inc_a.sum()/1e8:.2f} 亿", "说明": f"规模增幅 {inc_a.sum()/cur_amt*100:.1f}%"},
    {"模块": "方案A(达标5万)", "指标": "整体fpd7_rate$", "数值": f"{fpd7s_a*100:.2f}%",
     "说明": f"现状 {cur_fpd7*100:.2f}%, 降低 {(cur_fpd7-fpd7s_a)*100:.2f}pp (相对{(cur_fpd7-fpd7s_a)/cur_fpd7*100:.1f}%)"},
    {"模块": "方案B(满档位)", "指标": "可达平均额度", "数值": f"{avg_b:,.0f} 元",
     "说明": f"新增 {inc_b.sum()/1e8:.2f} 亿, fpd7_rate$→{fpd7s_b*100:.2f}%"},
])

# ---- 8. 输出 Excel ----
wb = load_workbook(SRC)
for s in list(wb.sheetnames):
    wb.remove(wb[s])
HDR_FILL = PatternFill("solid", fgColor="4E83FD")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=9)
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_sheet(ws, pct_cols=(), money_cols=(), int_cols=(), scale_cols=()):
    ncol = ws.max_column
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 24
    hdrs = [str(ws.cell(row=1, column=c).value or "") for c in range(1, ncol + 1)]
    for c, h in enumerate(hdrs, 1):
        col = get_column_letter(c)
        ws.column_dimensions[col].width = max(10, min(36, len(h) * 2 + 4))
        if h in pct_cols:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=c).number_format = "0.0%"
        elif h in money_cols:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=c).number_format = "#,##0"
        elif h in int_cols:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=c).number_format = "#,##0"
        if h in scale_cols and ws.max_row > 2:
            rng = f"{col}2:{col}{ws.max_row}"
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50,
                mid_color="FFEB84", end_type="max", end_color="F8696B"))
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{ws.max_row}"

def write_df(ws, data):
    for c, col in enumerate(data.columns, 1):
        ws.cell(row=1, column=c, value=col)
    for r, (_, row) in enumerate(data.iterrows(), 2):
        for c, col in enumerate(data.columns, 1):
            ws.cell(row=r, column=c, value=row[col])

# 原始数据
wb_src = load_workbook(SRC)
ws5 = wb.create_sheet("Sheet3_原始")
for row in wb_src.worksheets[0].iter_rows():
    for cell in row:
        if cell.value is None and not cell.has_style:
            continue
        nc = ws5.cell(row=cell.row, column=cell.column, value=cell.value)
        if cell.has_style:
            nc._style = cell._style
wb_src.close()

# 提额思路
ws_i = wb.create_sheet("提额思路")
write_df(ws_i, idea)
for c in range(1, 3):
    cell = ws_i.cell(row=1, column=c)
    cell.fill, cell.font = HDR_FILL, HDR_FONT
ws_i.freeze_panes = "A2"
ws_i.column_dimensions["A"].width = 16
ws_i.column_dimensions["B"].width = 95
for r in range(2, ws_i.max_row + 1):
    ws_i.cell(row=r, column=1).font = Font(name="微软雅黑", size=9, bold=True)
    ws_i.cell(row=r, column=2).font = BODY_FONT
    for c in range(1, 3):
        ws_i.cell(row=r, column=c).border = BORDER

# 提幅矩阵
ws_m = wb.create_sheet("提幅矩阵")
write_df(ws_m, matrix_df)
style_sheet(ws_m, pct_cols=[], money_cols=[], int_cols=["基础档%"])
for r in range(2, ws_m.max_row + 1):
    for c in range(3, ws_m.max_column + 1):
        cell = ws_m.cell(row=r, column=c)
        if isinstance(cell.value, (int, float)):
            cell.number_format = "0%"
            if cell.value >= 2.0:
                cell.fill = PatternFill("solid", fgColor="FDE8E8")
            elif cell.value >= 1.0:
                cell.fill = PatternFill("solid", fgColor="FFF3CD")
            else:
                cell.fill = PatternFill("solid", fgColor="E7F6EE")

# 提额结果
res = df[["flag_customer", "use_rate_group", "借款次数", "评级", "结清客户", "额度使用率",
          "件均", "平均额度", "提额客户", "提额客户_平均额度", "提额客户_提额后平均",
          "未提额客户", "未提额_额度件均", "未提额客户_提后件均", "发起下一笔订单",
          "fpd1_fm_dd", "fpd7_fm_dd", "fpd1_rate", "fpd1_rate$", "fpd7_rate", "fpd7_rate$",
          "系数盖帽", "额度盖帽", "风险档", "风险系数", "使用率基础", "是否可提额",
          "方案A幅度", "方案A后额度", "方案A新增(万)", "方案B幅度", "方案B后额度", "提额决策依据"]].copy()
ws = wb.create_sheet("提额结果")
write_df(ws, res)
style_sheet(ws, pct_cols=["方案A幅度", "方案B幅度", "风险系数", "使用率基础", "fpd1_rate", "fpd1_rate$", "fpd7_rate", "fpd7_rate$", "系数盖帽"],
            money_cols=["平均额度", "件均", "提额客户_平均额度", "提额客户_提额后平均", "未提额_额度件均",
                        "未提额客户_提后件均", "额度盖帽", "方案A后额度", "方案B后额度"],
            int_cols=["结清客户", "提额客户", "未提额客户", "发起下一笔订单", "fpd1_fm_dd", "fpd7_fm_dd", "方案A新增(万)"],
            scale_cols=["fpd7_rate$", "方案A幅度", "方案A后额度"])

# 分档统计
tier_recs = []
for rn, g in df.groupby("风险档", observed=True, sort=False):
    tier_recs.append({
        "风险档": rn, "客群数": len(g),
        "结清客户数": int(g["结清客户"].sum()),
        "提额客户数": int(g["提额客户"].sum()),
        "风险系数": g["风险系数"].iloc[0],
        "方案A平均幅度%": round((g["提额客户"] * g["方案A幅度"]).sum() / g["提额客户"].sum() * 100, 0) if g["提额客户"].sum() else 0,
        "方案A新增(万)": round(g["方案A新增"].sum() / 1e4, 0),
        "fpd7_rate$加权%": round((g["结清客户"] * g["平均额度"] * g["fpd7_rate$"]).sum() /
                                  (g["结清客户"] * g["平均额度"]).sum() * 100, 2),
    })
tier_stat = pd.DataFrame(tier_recs)
tier_stat["风险档"] = pd.Categorical(tier_stat["风险档"], categories=order, ordered=True)
tier_stat = tier_stat.sort_values("风险档")
ws2 = wb.create_sheet("分档统计")
write_df(ws2, tier_stat)
style_sheet(ws2, pct_cols=["风险系数", "方案A平均幅度%", "fpd7_rate$加权%"],
            int_cols=["结清客户数", "提额客户数", "方案A新增(万)"], scale_cols=["fpd7_rate$加权%", "方案A新增(万)"])

# 方案总结
ws6 = wb.create_sheet("提额方案总结")
write_df(ws6, summary)
for c in range(1, 5):
    cell = ws6.cell(row=1, column=c)
    cell.fill, cell.font = HDR_FILL, HDR_FONT
ws6.freeze_panes = "A2"
for col, wd in zip("ABCD", [18, 26, 40, 50]):
    ws6.column_dimensions[col].width = wd
for r in range(2, ws6.max_row + 1):
    for c in range(1, 5):
        cell = ws6.cell(row=r, column=c)
        cell.font = BODY_FONT
        cell.border = BORDER

wb.save(DST)
print(f"\n✅ 已生成: {DST}")
print("Sheets:", wb.sheetnames)
