from __future__ import annotations

"""Build a row-level strategy execution workbook from the coefficient table.

Every row of 提额系数20260812.xlsx is retained.  The script appends read-only
online outcome statistics for the matching 2026-08-12 new-strategy cohort.
"""

import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analyze_deployed_coefficient import lookup_key, normalized_code, policy_lookup
from probe_flexi_cash_strategy import get_odps, query

TABLE = "pb_biz_credit.flexi_cash_jq_result_v1"
COEFFICIENT_FILE = Path("C:/Users/PP-2026070302/Desktop") / "提额系数20260812.xlsx"
OUTPUT = Path("risk_analysis/提额策略执行分析_20260812_v2.xlsx")
JSON_OUTPUT = Path("risk_analysis/strategy_execution_by_cell_20260812.json")
TARGET_AVERAGE_YUAN = 50_000.0


def value_as_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def zero(value: Any) -> float:
    return value_as_float(value) or 0.0


def safe_div(numerator: float, denominator: float) -> float | None:
    return numerator / denominator if denominator else None


def rate_label(value: float | None) -> str:
    if value is None:
        return "空"
    return f"{value:g}"


def new_metric() -> dict[str, Any]:
    return {
        "total_cnt": 0.0,
        "raised_cnt": 0.0,
        "non_raised_cnt": 0.0,
        "other_te_cnt": 0.0,
        "actual_raise_cnt": 0.0,
        "actual_unchanged_cnt": 0.0,
        "actual_decrease_cnt": 0.0,
        "raised_actual_raise_cnt": 0.0,
        "non_raised_actual_unchanged_cnt": 0.0,
        "raised_before_sum_cent": 0.0,
        "raised_after_sum_cent": 0.0,
        "raised_increment_sum_cent": 0.0,
        "non_raised_before_sum_cent": 0.0,
        "non_raised_after_sum_cent": 0.0,
        "raised_t0_cnt": 0.0,
        "non_raised_t0_cnt": 0.0,
        "raised_next_loan_cnt": 0.0,
        "non_raised_next_loan_cnt": 0.0,
        "raised_fpd1_num": 0.0,
        "raised_fpd1_den": 0.0,
        "non_raised_fpd1_num": 0.0,
        "non_raised_fpd1_den": 0.0,
        "raised_fpd7_num": 0.0,
        "raised_fpd7_den": 0.0,
        "non_raised_fpd7_num": 0.0,
        "non_raised_fpd7_den": 0.0,
        "raised_fpd15_num": 0.0,
        "raised_fpd15_den": 0.0,
        "non_raised_fpd15_num": 0.0,
        "non_raised_fpd15_den": 0.0,
        "raised_fpd30_num": 0.0,
        "raised_fpd30_den": 0.0,
        "non_raised_fpd30_num": 0.0,
        "non_raised_fpd30_den": 0.0,
        "coefficient_match_cnt": 0.0,
        "coefficient_mismatch_cnt": 0.0,
        "positive_config_non_raised_cnt": 0.0,
        "zero_config_raised_cnt": 0.0,
        "raised_configured_rate_distribution": defaultdict(float),
    }


def add_group(metric: dict[str, Any], row: dict[str, Any], expected_rate: float | None) -> None:
    """Add one SQL aggregate row into a cell/overall metric bucket."""
    cnt = zero(row["cnt"])
    flag = normalized_code(row["te_flag"])
    actual_rate = value_as_float(row["cash_remark1"])
    prefix = "raised" if flag == "1" else "non_raised" if flag == "0" else "other"

    metric["total_cnt"] += cnt
    metric["actual_raise_cnt"] += zero(row["actual_raise_cnt"])
    metric["actual_unchanged_cnt"] += zero(row["actual_unchanged_cnt"])
    metric["actual_decrease_cnt"] += zero(row["actual_decrease_cnt"])

    # Coefficient comparison is intentionally limited to eligible/raised
    # customers (te_flag=1), per the strategy-analysis definition.
    if expected_rate is not None and prefix == "raised":
        distribution_key = rate_label(actual_rate)
        metric["raised_configured_rate_distribution"][distribution_key] += cnt
        if actual_rate is not None and abs(actual_rate - expected_rate) < 1e-10:
            metric["coefficient_match_cnt"] += cnt
        else:
            metric["coefficient_mismatch_cnt"] += cnt

    if prefix == "other":
        metric["other_te_cnt"] += cnt
        return

    metric[f"{prefix}_cnt"] += cnt
    metric[f"{prefix}_before_sum_cent"] += zero(row["before_sum_cent"])
    metric[f"{prefix}_after_sum_cent"] += zero(row["after_sum_cent"])
    metric[f"{prefix}_t0_cnt"] += zero(row["t0_cnt"])
    metric[f"{prefix}_next_loan_cnt"] += zero(row["next_loan_cnt"])
    for horizon in ("fpd1", "fpd7", "fpd15", "fpd30"):
        metric[f"{prefix}_{horizon}_num"] += zero(row[f"{horizon}_num"])
        metric[f"{prefix}_{horizon}_den"] += zero(row[f"{horizon}_den"])

    if prefix == "raised":
        metric["raised_increment_sum_cent"] += zero(row["increment_sum_cent"])
        metric["raised_actual_raise_cnt"] += zero(row["actual_raise_cnt"])
        if actual_rate is None or abs(actual_rate) < 1e-10:
            metric["zero_config_raised_cnt"] += zero(row["actual_raise_cnt"])
    else:
        metric["non_raised_actual_unchanged_cnt"] += zero(row["actual_unchanged_cnt"])
        if actual_rate is not None and actual_rate > 0:
            metric["positive_config_non_raised_cnt"] += cnt


def distribution_label(distribution: dict[str, float]) -> str:
    def sort_key(item: tuple[str, float]) -> tuple[int, float | str]:
        try:
            return (0, float(item[0]))
        except ValueError:
            return (1, item[0])

    return "；".join(f"{rate}:{int(count):,}" for rate, count in sorted(distribution.items(), key=sort_key)) or "—"


def calculate(metric: dict[str, Any], expected_rate: float | None = None) -> dict[str, Any]:
    raised_count = metric["raised_cnt"]
    non_raised_count = metric["non_raised_cnt"]
    total_count = metric["total_cnt"]
    settled_count = raised_count + non_raised_count
    before_sum = metric["raised_before_sum_cent"] + metric["non_raised_before_sum_cent"]
    after_sum = metric["raised_after_sum_cent"] + metric["non_raised_after_sum_cent"]
    config_total = metric["coefficient_match_cnt"] + metric["coefficient_mismatch_cnt"]
    execution_numerator = metric["raised_actual_raise_cnt"] + metric["non_raised_actual_unchanged_cnt"]

    result: dict[str, Any] = {
        "结清客户数": int(settled_count),
        "提额客户数": int(raised_count),
        "未提额客户数": int(non_raised_count),
        "其他标签客户数": int(metric["other_te_cnt"]),
        "提额覆盖率": safe_div(raised_count, settled_count),
        "满足提额条件客户线上配置系数分布(cash_remark1)": distribution_label(metric["raised_configured_rate_distribution"]),
        "满足提额条件客户线上系数一致客户数": int(metric["coefficient_match_cnt"]),
        "满足提额条件客户线上系数不一致客户数": int(metric["coefficient_mismatch_cnt"]),
        "满足提额条件客户线上系数一致率": safe_div(metric["coefficient_match_cnt"], config_total),
        "提额客户提额前平均额度(元)": safe_div(metric["raised_before_sum_cent"], raised_count * 100),
        "提额客户提额后平均额度(元)": safe_div(metric["raised_after_sum_cent"], raised_count * 100),
        "提额客户实际提幅(加权)": safe_div(metric["raised_increment_sum_cent"], metric["raised_before_sum_cent"]),
        "提额客户户均增额(元)": safe_div(metric["raised_increment_sum_cent"], raised_count * 100),
        "未提额客户提额前平均额度(元)": safe_div(metric["non_raised_before_sum_cent"], non_raised_count * 100),
        "未提额客户提额后平均额度(元)": safe_div(metric["non_raised_after_sum_cent"], non_raised_count * 100),
        "结清客户提额前平均额度(元)": safe_div(before_sum, settled_count * 100),
        "结清客户提额后平均额度(元)": safe_div(after_sum, settled_count * 100),
        "结清客户额度提升率": safe_div(after_sum - before_sum, before_sum),
        "提额客户T0发起率": safe_div(metric["raised_t0_cnt"], raised_count),
        "未提额客户T0发起率": safe_div(metric["non_raised_t0_cnt"], non_raised_count),
        "提额客户下一笔发起率": safe_div(metric["raised_next_loan_cnt"], raised_count),
        "未提额客户下一笔发起率": safe_div(metric["non_raised_next_loan_cnt"], non_raised_count),
        "提额客户实际提额率": safe_div(metric["raised_actual_raise_cnt"], raised_count),
        "未提额客户实际未变率": safe_div(metric["non_raised_actual_unchanged_cnt"], non_raised_count),
        "提额/未提额执行符合率": safe_div(execution_numerator, settled_count),
        "正系数但未提额客户数": int(metric["positive_config_non_raised_cnt"]),
        "零系数但实际提额客户数": int(metric["zero_config_raised_cnt"]),
        "实际提额人数": int(metric["actual_raise_cnt"]),
        "实际未变人数": int(metric["actual_unchanged_cnt"]),
        "实际降额人数": int(metric["actual_decrease_cnt"]),
    }
    for horizon in ("FPD1", "FPD7", "FPD15", "FPD30"):
        lower = horizon.lower()
        result[f"提额客户{horizon}"] = safe_div(metric[f"raised_{lower}_num"], metric[f"raised_{lower}_den"])
        result[f"提额客户{horizon}样本"] = int(metric[f"raised_{lower}_den"])
        result[f"未提额客户{horizon}"] = safe_div(metric[f"non_raised_{lower}_num"], metric[f"non_raised_{lower}_den"])
        result[f"未提额客户{horizon}样本"] = int(metric[f"non_raised_{lower}_den"])

    execution_rate = result["提额/未提额执行符合率"]
    coefficient_rate = result["满足提额条件客户线上系数一致率"]
    if settled_count == 0:
        conclusion = "本批次无样本"
    elif execution_rate is not None and execution_rate < 1:
        conclusion = "提额/未提额执行异常"
    elif coefficient_rate is None:
        conclusion = "无法映射系数表"
    elif coefficient_rate < 1:
        conclusion = "提额执行符合；系数需核对"
    else:
        conclusion = "系数与提额执行均符合"
    result["策略执行结论"] = conclusion
    if expected_rate is not None:
        result["表内目标提额系数"] = expected_rate
    return result


def add_overall_attainment(summary: dict[str, Any]) -> dict[str, Any]:
    result = dict(summary)
    count = result["结清客户数"]
    before_avg = result["结清客户提额前平均额度(元)"]
    after_avg = result["结清客户提额后平均额度(元)"]
    result.update(
        {
            "目标户均额度(元)": TARGET_AVERAGE_YUAN,
            "提额前目标达成率": safe_div(before_avg or 0, TARGET_AVERAGE_YUAN),
            "提额后目标达成率": safe_div(after_avg or 0, TARGET_AVERAGE_YUAN),
            "提额后户均差距(元)": TARGET_AVERAGE_YUAN - (after_avg or 0),
            "实际总增额(元)": ((after_avg or 0) - (before_avg or 0)) * count,
            "达标所需总增额(元)": (TARGET_AVERAGE_YUAN - (before_avg or 0)) * count,
        }
    )
    result["目标增量达成率"] = safe_div(result["实际总增额(元)"], result["达标所需总增额(元)"])
    return result


def main() -> None:
    policy = policy_lookup()
    odps = get_odps()
    sql = f"""
      SELECT
        jq_customer_flag, defq_use_rate_level, separate_installment_loan_cnt,
        jq_mix_score_bin, cash_remark1, te_flag,
        COUNT(1) AS cnt,
        SUM(CAST(jq_credit_quota AS DOUBLE)) AS before_sum_cent,
        SUM(CAST(cash_quota_amount_now_after AS DOUBLE)) AS after_sum_cent,
        SUM(CAST(cash_quota_amount_now_after AS DOUBLE) - CAST(jq_credit_quota AS DOUBLE)) AS increment_sum_cent,
        SUM(CASE WHEN CAST(cash_quota_amount_now_after AS DOUBLE) > CAST(jq_credit_quota AS DOUBLE) THEN 1 ELSE 0 END) AS actual_raise_cnt,
        SUM(CASE WHEN CAST(cash_quota_amount_now_after AS DOUBLE) = CAST(jq_credit_quota AS DOUBLE) THEN 1 ELSE 0 END) AS actual_unchanged_cnt,
        SUM(CASE WHEN CAST(cash_quota_amount_now_after AS DOUBLE) < CAST(jq_credit_quota AS DOUBLE) THEN 1 ELSE 0 END) AS actual_decrease_cnt,
        SUM(CASE WHEN CAST(next_curday_flag AS BIGINT) = 1 THEN 1 ELSE 0 END) AS t0_cnt,
        SUM(CASE WHEN CAST(next_loan_flag AS BIGINT) = 1 THEN 1 ELSE 0 END) AS next_loan_cnt,
        SUM(CAST(fpd1_fz_dd AS DOUBLE)) AS fpd1_num,
        SUM(CAST(fpd1_fm_dd AS DOUBLE)) AS fpd1_den,
        SUM(CAST(fpd7_fz_dd AS DOUBLE)) AS fpd7_num,
        SUM(CAST(fpd7_fm_dd AS DOUBLE)) AS fpd7_den,
        SUM(CAST(fpd15_fz_dd AS DOUBLE)) AS fpd15_num,
        SUM(CAST(fpd15_fm_dd AS DOUBLE)) AS fpd15_den,
        SUM(CAST(fpd30_fz_dd AS DOUBLE)) AS fpd30_num,
        SUM(CAST(fpd30_fm_dd AS DOUBLE)) AS fpd30_den
      FROM {TABLE}
      WHERE jq_date = '2026-08-12'
        AND str_type = 'new'
        AND CAST(jq_credit_quota AS DOUBLE) >= 0
        AND CAST(cash_quota_amount_now_after AS DOUBLE) >= 0
      GROUP BY jq_customer_flag, defq_use_rate_level, separate_installment_loan_cnt,
               jq_mix_score_bin, cash_remark1, te_flag
    """
    grouped_rows = query(odps, sql)

    metrics_by_key: dict[tuple[str, str, int, str], dict[str, Any]] = defaultdict(new_metric)
    full_metrics = new_metric()
    mapped_metrics = new_metric()
    unmapped_metrics = new_metric()

    for row in grouped_rows:
        key = lookup_key(
            row["jq_customer_flag"], row["defq_use_rate_level"],
            row["separate_installment_loan_cnt"], row["jq_mix_score_bin"],
        )
        expected_rate = policy.get(key) if key is not None else None
        add_group(full_metrics, row, expected_rate)
        if expected_rate is None or key is None:
            add_group(unmapped_metrics, row, None)
            continue
        add_group(mapped_metrics, row, expected_rate)
        add_group(metrics_by_key[key], row, expected_rate)

    source_wb = load_workbook(COEFFICIENT_FILE, data_only=True, read_only=True)
    source_ws = source_wb[source_wb.sheetnames[0]]
    policy_rows = [row for row in source_ws.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row)]

    detailed_rows: list[dict[str, Any]] = []
    for source_row in policy_rows:
        key = (str(source_row[0]).strip(), str(source_row[1]).strip(), int(float(source_row[2])), str(source_row[3]).strip())
        expected_rate = float(source_row[4])
        detail = calculate(metrics_by_key.get(key, new_metric()), expected_rate)
        detail["策略格"] = " | ".join(map(str, key))
        detailed_rows.append(detail)

    full_summary = add_overall_attainment(calculate(full_metrics))
    mapped_summary = add_overall_attainment(calculate(mapped_metrics))
    unmapped_summary = calculate(unmapped_metrics)

    # Preserve the original table exactly, then append business execution fields.
    workbook = load_workbook(COEFFICIENT_FILE)
    sheet = workbook[workbook.sheetnames[0]]
    sheet.title = "策略执行分析"
    appended_headers = [
        "结清客户数", "提额客户数", "未提额客户数", "提额覆盖率",
        "满足提额条件客户线上配置系数分布(cash_remark1)", "满足提额条件客户线上系数一致客户数", "满足提额条件客户线上系数不一致客户数", "满足提额条件客户线上系数一致率",
        "提额客户提额前平均额度(元)", "提额客户提额后平均额度(元)", "提额客户实际提幅(加权)", "提额客户户均增额(元)",
        "未提额客户提额前平均额度(元)", "未提额客户提额后平均额度(元)",
        "结清客户提额前平均额度(元)", "结清客户提额后平均额度(元)", "结清客户额度提升率",
        "提额客户T0发起率", "未提额客户T0发起率", "提额客户下一笔发起率", "未提额客户下一笔发起率",
        "提额客户FPD1", "提额客户FPD1样本", "未提额客户FPD1", "未提额客户FPD1样本",
        "提额客户FPD7", "提额客户FPD7样本", "未提额客户FPD7", "未提额客户FPD7样本",
        "提额客户FPD15", "提额客户FPD15样本", "未提额客户FPD15", "未提额客户FPD15样本",
        "提额客户FPD30", "提额客户FPD30样本", "未提额客户FPD30", "未提额客户FPD30样本",
        "提额客户实际提额率", "未提额客户实际未变率", "提额/未提额执行符合率",
        "正系数但未提额客户数", "零系数但实际提额客户数", "策略执行结论",
    ]
    first_appended_column = sheet.max_column + 1
    for column, header in enumerate(appended_headers, first_appended_column):
        sheet.cell(1, column, header)
    for row_index, detail in enumerate(detailed_rows, 2):
        for column, header in enumerate(appended_headers, first_appended_column):
            sheet.cell(row_index, column, detail[header])

    navy = "1F4E78"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    original_widths = [18, 18, 11, 11, 16]
    for column, width in enumerate(original_widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(first_appended_column, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 17
    distribution_column = first_appended_column + appended_headers.index("满足提额条件客户线上配置系数分布(cash_remark1)")
    conclusion_column = first_appended_column + appended_headers.index("策略执行结论")
    sheet.column_dimensions[get_column_letter(distribution_column)].width = 25
    sheet.column_dimensions[get_column_letter(conclusion_column)].width = 26

    percentage_headers = {
        "提额覆盖率", "满足提额条件客户线上系数一致率", "提额客户实际提幅(加权)", "结清客户额度提升率",
        "提额客户T0发起率", "未提额客户T0发起率", "提额客户下一笔发起率", "未提额客户下一笔发起率",
        "提额客户FPD1", "未提额客户FPD1", "提额客户FPD7", "未提额客户FPD7",
        "提额客户FPD15", "未提额客户FPD15", "提额客户FPD30", "未提额客户FPD30",
        "提额客户实际提额率", "未提额客户实际未变率", "提额/未提额执行符合率",
    }
    amount_headers = {
        "提额客户提额前平均额度(元)", "提额客户提额后平均额度(元)", "提额客户户均增额(元)",
        "未提额客户提额前平均额度(元)", "未提额客户提额后平均额度(元)",
        "结清客户提额前平均额度(元)", "结清客户提额后平均额度(元)",
    }
    for header in percentage_headers:
        column = first_appended_column + appended_headers.index(header)
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row_index, column).number_format = "0.00%"
    for header in amount_headers:
        column = first_appended_column + appended_headers.index(header)
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row_index, column).number_format = "#,##0.00"
    match_rate_column = first_appended_column + appended_headers.index("满足提额条件客户线上系数一致率")
    sheet.conditional_formatting.add(
        f"{get_column_letter(match_rate_column)}2:{get_column_letter(match_rate_column)}{sheet.max_row}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, conclusion_column).alignment = Alignment(wrap_text=True, vertical="center")

    overview = workbook.create_sheet("整体达成")
    overview.append(["提额策略整体达成", "结果", "说明"])
    overview_rows = [
        ("统计范围", "2026-08-12 · str_type=new", "提额前/后额度有效的结清样本。"),
        ("目标户均额度", TARGET_AVERAGE_YUAN, "沿用 5 万元户均额度目标。"),
        ("全量有效结清样本", full_summary["结清客户数"], "包含 1 条无法映射至315格策略表的异常维度样本。"),
        ("策略表可映射结清样本", mapped_summary["结清客户数"], "逐格统计表的合计口径。"),
        ("提额前户均额度", full_summary["结清客户提额前平均额度(元)"], "全量有效结清样本。"),
        ("提额前目标达成率", full_summary["提额前目标达成率"], "提额前户均额度 / 50,000。"),
        ("提额后户均额度", full_summary["结清客户提额后平均额度(元)"], "全量有效结清样本。"),
        ("提额后目标达成率", full_summary["提额后目标达成率"], "提额后户均额度 / 50,000。"),
        ("提额后距目标差距", full_summary["提额后户均差距(元)"], "每个结清样本距 50,000 元的平均差距。"),
        ("实际总增额", full_summary["实际总增额(元)"], "提额后总额度 - 提额前总额度。"),
        ("达标所需总增额", full_summary["达标所需总增额(元)"], "若全体户均达到 50,000 元所需增额。"),
        ("目标增量达成率", full_summary["目标增量达成率"], "实际总增额 / 达标所需总增额。"),
        ("提额客户数", full_summary["提额客户数"], "te_flag=1。"),
        ("提额覆盖率", full_summary["提额覆盖率"], "提额客户数 / 全量有效结清样本。"),
        ("提额客户提额前户均额度", full_summary["提额客户提额前平均额度(元)"], "te_flag=1。"),
        ("提额客户提额后户均额度", full_summary["提额客户提额后平均额度(元)"], "te_flag=1。"),
        ("提额客户实际提幅", full_summary["提额客户实际提幅(加权)"], "提额客户总增额 / 提额前总额度。"),
        ("未提额客户数", full_summary["未提额客户数"], "te_flag=0。"),
        ("未提额客户提额前户均额度", full_summary["未提额客户提额前平均额度(元)"], "te_flag=0。"),
        ("未提额客户提额后户均额度", full_summary["未提额客户提额后平均额度(元)"], "te_flag=0。"),
        ("满足提额条件客户线上系数一致率", full_summary["满足提额条件客户线上系数一致率"], "仅 te_flag=1 客户；cash_remark1 与策略表目标系数一致的样本占比。"),
        ("提额/未提额执行符合率", full_summary["提额/未提额执行符合率"], "te_flag=1 实际提额，te_flag=0 实际额度不变。"),
        ("正系数但未提额客户", full_summary["正系数但未提额客户数"], "需结合资格、拦截或封顶原因码判断是否为策略例外。"),
        ("零系数但实际提额客户", full_summary["零系数但实际提额客户数"], "需核对公式例外或 cash_remark1 字段语义。"),
    ]
    for row in overview_rows:
        overview.append(row)

    overview.append([])
    overview.append(["贷后与发起表现", "提额客户", "未提额客户"])
    for horizon in ("T0发起率", "下一笔发起率", "FPD1", "FPD7", "FPD15", "FPD30"):
        if horizon.startswith("FPD"):
            overview.append((
                horizon,
                full_summary[f"提额客户{horizon}"],
                full_summary[f"未提额客户{horizon}"],
            ))
            overview.append((
                f"{horizon}样本数",
                full_summary[f"提额客户{horizon}样本"],
                full_summary[f"未提额客户{horizon}样本"],
            ))
        else:
            overview.append((horizon, full_summary[f"提额客户{horizon}"], full_summary[f"未提额客户{horizon}"]))

    for cell in overview[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    section_row = len(overview_rows) + 3
    for cell in overview[section_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3772F6")
        cell.alignment = Alignment(horizontal="center")
    overview.column_dimensions["A"].width = 33
    overview.column_dimensions["B"].width = 25
    overview.column_dimensions["C"].width = 70
    overview.freeze_panes = "A2"
    for row_index in range(2, overview.max_row + 1):
        overview.cell(row_index, 2).alignment = Alignment(vertical="center")
        overview.cell(row_index, 3).alignment = Alignment(wrap_text=True, vertical="center")
    percentage_summary_rows = {7, 9, 13, 15, 18, 22, 23, section_row + 1, section_row + 2, section_row + 3, section_row + 5, section_row + 7, section_row + 9, section_row + 11}
    for row_index in percentage_summary_rows:
        if row_index <= overview.max_row:
            overview.cell(row_index, 2).number_format = "0.00%"
            if row_index >= section_row:
                overview.cell(row_index, 3).number_format = "0.00%"
    amount_summary_rows = {3, 6, 8, 10, 11, 16, 17, 20, 21}
    for row_index in amount_summary_rows:
        if row_index <= overview.max_row:
            overview.cell(row_index, 2).number_format = "#,##0.00"

    method = workbook.create_sheet("口径说明")
    method.append(["项目", "说明"])
    notes = [
        ("行级统计方式", "保留提额系数表全部315行；在线上结果表按客户类型、额度使用率、借款数、风险评级映射并聚合。"),
        ("借款数映射", "策略表借款数 1–6 各自匹配；策略格 7 表示借款数 7 及以上。"),
        ("线上配置系数分布", "字段为 cash_remark1（线上配置/结果中的提额幅度）；仅展示 te_flag=1、即满足提额条件客户实际出现的系数及样本数。"),
        ("线上系数比对", "仅在 te_flag=1（满足提额条件）客户中，将 cash_remark1 与该策略行的目标提额系数直接比较；未提额客户不参与系数一致客户数、不一致客户数或一致率。"),
        ("提额/未提额执行", "te_flag=1 且实际提额、te_flag=0 且额度不变，计为执行符合。该口径验证结果是否按标签执行。"),
        ("正系数但未提额", "不自动判定为策略错误；若存在资格、拦截、额度封顶等额外闸门，需要结合原因码复核。"),
        ("T0发起率", "next_curday_flag=1 / 对应提额或未提额样本数。"),
        ("下一笔发起率", "next_loan_flag=1 / 对应提额或未提额样本数。"),
        ("FPD口径", "客户口径：fpdX_fz_dd / fpdX_fm_dd。每个FPD指标旁保留分母样本数；成熟样本较少时不可仅比较比率。"),
        ("金额单位", "原始额度字段单位为分；本工作簿新增平均额度、总增额字段均已转换为元。"),
        ("隐私与凭证", "仅输出逐格聚合结果，不导出客户标识，也不在工作簿中保留 ODPS 凭证。"),
    ]
    for note in notes:
        method.append(note)
    for cell in method[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
    method.column_dimensions["A"].width = 24
    method.column_dimensions["B"].width = 115
    for row in method.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    JSON_OUTPUT.write_text(
        json.dumps(
            {
                "scope": {
                    "business_date": "2026-08-12",
                    "strategy_scope": "str_type=new",
                    "table": TABLE,
                    "target_average_yuan": TARGET_AVERAGE_YUAN,
                    "loan_count_rule": "策略格7表示借款数7及以上",
                },
                "overall": {
                    "full_valid": full_summary,
                    "mapped_strategy_table": mapped_summary,
                    "unmapped": unmapped_summary,
                },
                "rows": detailed_rows,
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT}")
    print(f"Wrote {JSON_OUTPUT}")


if __name__ == "__main__":
    main()
