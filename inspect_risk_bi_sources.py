from __future__ import annotations

import ast
import json
import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

BASE = Path(r"D:\vscode\归因分析")
OUT = Path(r"E:\agent\monitor\_source_inspection")
OUT.mkdir(parents=True, exist_ok=True)

DOCX = BASE / "贷前申请量异常归因监控方案_v2.2_去除超额申请量升级逻辑.docx"
RESULT_XLSX = BASE / "贷前申请量异常归因结果_v2.2_20260612-20260622样本_分窗口配色.xlsx"
CONFIG_XLSX = BASE / "配置文档.xlsx"
NOTEBOOK = Path(r"D:\vscode\数据库连接示例.ipynb")


def color_desc(cell: Any) -> str:
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return ""
    fg = fill.fgColor
    value = getattr(fg, "rgb", None) or getattr(fg, "indexed", None) or getattr(fg, "theme", None) or ""
    return f"{fill.fill_type}:{fg.type}:{value}"


def write_docx_report() -> None:
    document = Document(DOCX)
    lines: list[str] = []
    lines.append(f"# {DOCX.name}")
    lines.append("")
    lines.append(f"段落数：{len(document.paragraphs)}；表格数：{len(document.tables)}")
    lines.append("")
    lines.append("## 段落")
    for i, p in enumerate(document.paragraphs, 1):
        text = p.text.strip()
        if text:
            lines.append(f"[{i:03d}] ({p.style.name}) {text}")
    lines.append("")
    lines.append("## 表格")
    for ti, table in enumerate(document.tables, 1):
        lines.append(f"### 表格 {ti}（{len(table.rows)} 行 × {len(table.columns)} 列）")
        for ri, row in enumerate(table.rows, 1):
            values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            lines.append(f"R{ri}: " + " | ".join(values))
        lines.append("")
    with zipfile.ZipFile(DOCX) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    lines.append("## 内嵌图片")
    lines.extend(f"- {name}" for name in media) if media else lines.append("- 无")
    (OUT / "implementation_plan.md").write_text("\n".join(lines), encoding="utf-8")


def render_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " / ").replace("\r", " ")
    return text[:500]


def write_workbook_report(path: Path, output_name: str) -> None:
    wb = load_workbook(path, data_only=False, read_only=False)
    lines: list[str] = [f"# {path.name}", ""]
    lines.append("工作表：" + "、".join(wb.sheetnames))
    for ws in wb.worksheets:
        lines.append("")
        lines.append(f"## {ws.title}")
        lines.append(f"维度：{ws.max_row} 行 × {ws.max_column} 列；冻结窗格：{ws.freeze_panes or '无'}")
        if ws.merged_cells.ranges:
            lines.append("合并单元格：" + ", ".join(str(x) for x in ws.merged_cells.ranges))
        cf = list(ws.conditional_formatting)
        if cf:
            lines.append("条件格式范围：" + ", ".join(str(x) for x in cf))
        lines.append("")
        lines.append("### 非空单元格（坐标 / 值 / 填充色）")
        non_empty = 0
        for row in ws.iter_rows():
            cells: list[str] = []
            for cell in row:
                if cell.value is not None:
                    non_empty += 1
                    fill = color_desc(cell)
                    suffix = f" [fill={fill}]" if fill else ""
                    cells.append(f"{cell.coordinate}={render_value(cell.value)}{suffix}")
            if cells:
                lines.append(" | ".join(cells))
        lines.append(f"\n非空单元格数量：{non_empty}")
    (OUT / output_name).write_text("\n".join(lines), encoding="utf-8")


SENSITIVE_NAME = re.compile(r"(?:access.?key|secret|password|token|api.?key|^ak$|^sk$|credential)", re.I)


def target_name(node: ast.AST) -> str | None:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        return node.attr
    return None


def call_name(node: ast.AST) -> str | None:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        return node.attr
    return None


def write_notebook_report() -> None:
    nb = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    lines: list[str] = [f"# {NOTEBOOK.name}", "", f"单元格数：{len(nb.get('cells', []))}", ""]
    imports: set[str] = set()
    assignments: set[str] = set()
    endpoints: set[str] = set()
    odps_calls: list[str] = []
    sql_tables: set[str] = set()

    for idx, cell in enumerate(nb.get("cells", []), 1):
        if cell.get("cell_type") != "code":
            continue
        source = "".join(cell.get("source", []))
        try:
            tree = ast.parse(source)
        except SyntaxError:
            lines.append(f"- Code cell {idx}: 无法静态解析（已跳过源码展示，防止泄露凭据）")
            continue
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imports.update(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom):
                imports.add(f"{node.module or ''}." + ",".join(alias.name for alias in node.names))
            elif isinstance(node, (ast.Assign, ast.AnnAssign)):
                targets = node.targets if isinstance(node, ast.Assign) else [node.target]
                for target in targets:
                    name = target_name(target)
                    if name:
                        assignments.add(name)
                value = node.value
                if isinstance(value, ast.Constant) and isinstance(value.value, str) and "aliyun" in value.value.lower():
                    endpoints.add(value.value)
            elif isinstance(node, ast.Call):
                name = call_name(node.func)
                if name == "ODPS":
                    odps_calls.append(
                        f"- Code cell {idx}: ODPS(...)，位置参数={len(node.args)}，关键字参数={','.join((kw.arg or '**') for kw in node.keywords) or '无'}"
                    )
                elif name in {"execute_sql", "run_sql", "get_table", "read_table", "write_table"}:
                    odps_calls.append(f"- Code cell {idx}: 调用 {name}(...)" )
            elif isinstance(node, ast.Constant) and isinstance(node.value, str):
                for table in re.findall(r"\b[a-zA-Z_][\w]*\.[a-zA-Z_][\w]*\b", node.value):
                    sql_tables.add(table)

    lines.append("## 连接方式（已脱敏）")
    lines.append("导入：" + ("、".join(sorted(imports)) or "未识别"))
    safe_names = sorted(name for name in assignments if not SENSITIVE_NAME.search(name))
    secret_names = sorted(name for name in assignments if SENSITIVE_NAME.search(name))
    lines.append("非敏感变量：" + ("、".join(safe_names) or "未识别"))
    lines.append("凭据变量（仅名称，值未输出）：" + ("、".join(secret_names) or "未识别"))
    lines.append("Endpoint：" + ("、".join(sorted(endpoints)) or "未识别"))
    lines.extend(odps_calls or ["未识别到 ODPS 调用"])
    lines.append("SQL 中识别的表：" + ("、".join(sorted(sql_tables)) or "无"))
    (OUT / "notebook_connection_redacted.md").write_text("\n".join(lines), encoding="utf-8")


write_docx_report()
write_workbook_report(RESULT_XLSX, "attribution_result_workbook.md")
write_workbook_report(CONFIG_XLSX, "configuration_workbook.md")
write_notebook_report()
print(OUT)
