from __future__ import annotations

"""Aggregate post-launch impact / target-gap analysis for the flexi-cash strategy.

Amounts in the result table are in cents. Output metrics are converted to yuan
where named with `_yuan` and are aggregated only.
"""

import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from analyze_deployed_coefficient import lookup_key, normalized_code, policy_lookup
from probe_flexi_cash_strategy import get_odps, query

TABLE = "pb_biz_credit.flexi_cash_jq_result_v1"
DESKTOP = Path("C:/Users/PP-2026070302/Desktop")
PLAN_FILE = DESKTOP / "\u5927\u6a21\u578b\u6837\u672c2-\u63d0\u989d\u65b9\u6848-ABC\u5ba2\u7fa4-20260812.xlsx"
OUTPUT = Path("risk_analysis/strategy_impact_0812.json")


def as_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def plan_rates() -> dict[tuple[str, str, int, str], dict[str, float | bool]]:
    wb = load_workbook(PLAN_FILE, data_only=True, read_only=True)
    ws = wb["\u63d0\u989d\u7ed3\u679c"]
    result: dict[tuple[str, str, int, str], dict[str, float | bool]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        key = (str(row[0]).strip(), str(row[1]).strip(), int(float(row[2])), str(row[3]).strip())
        result[key] = {
            "plan_abc_eligible": bool(row[18]),
            "plan_a_rate": float(row[19] or 0),
            "plan_b_rate": float(row[22] or 0),
        }
    return result


def empty_metric() -> dict[str, float | int]:
    return {
        "cnt": 0,
        "before_sum_cent": 0.0,
        "after_sum_cent": 0.0,
        "actual_increment_cent": 0.0,
        "actual_rate_unconstrained_increment_cent": 0.0,
        "coefficient_table_increment_cent": 0.0,
        "plan_a_increment_cent": 0.0,
        "plan_b_increment_cent": 0.0,
        "positive_config_no_raise_cnt": 0,
        "coefficient_match_cnt": 0,
        "coefficient_mismatch_cnt": 0,
        "no_lookup_cnt": 0,
        "actual_rate_below_table_cnt": 0,
        "actual_rate_above_table_cnt": 0,
        "rate_table_shortfall_cent": 0.0,
        "plan_a_vs_table_gap_cent": 0.0,
    }


def enrich(metric: dict[str, float | int]) -> dict[str, float | int | None]:
    count = int(metric["cnt"])
    before = float(metric["before_sum_cent"])
    after = float(metric["after_sum_cent"])
    result = dict(metric)
    result.update(
        {
            "before_avg_yuan": before / count / 100 if count else None,
            "after_avg_yuan": after / count / 100 if count else None,
            "actual_increment_yuan": float(metric["actual_increment_cent"]) / 100,
            "actual_increment_per_customer_yuan": float(metric["actual_increment_cent"]) / count / 100 if count else None,
            "actual_raise_rate_by_amount": (after / before - 1) if before else None,
            "actual_rate_unconstrained_increment_yuan": float(metric["actual_rate_unconstrained_increment_cent"]) / 100,
            "coefficient_table_increment_yuan": float(metric["coefficient_table_increment_cent"]) / 100,
            "plan_a_increment_yuan": float(metric["plan_a_increment_cent"]) / 100,
            "plan_b_increment_yuan": float(metric["plan_b_increment_cent"]) / 100,
            "rate_table_shortfall_yuan": float(metric["rate_table_shortfall_cent"]) / 100,
            "plan_a_vs_table_gap_yuan": float(metric["plan_a_vs_table_gap_cent"]) / 100,
        }
    )
    return result


def main() -> None:
    odps = get_odps()
    summary_sql = f"""
      SELECT str_type, app_flow_flag, te_flag,
             COUNT(1) AS cnt,
             SUM(CAST(jq_credit_quota AS DOUBLE)) AS before_sum_cent,
             SUM(CAST(cash_quota_amount_now_after AS DOUBLE)) AS after_sum_cent,
             SUM(CAST(cash_quota_amount_now_after AS DOUBLE) - CAST(jq_credit_quota AS DOUBLE)) AS increment_cent,
             SUM(CASE WHEN CAST(cash_quota_amount_now_after AS DOUBLE) > CAST(jq_credit_quota AS DOUBLE) THEN 1 ELSE 0 END) AS actual_raise_cnt,
             SUM(CASE WHEN CAST(cash_quota_amount_now_after AS DOUBLE) = CAST(jq_credit_quota AS DOUBLE) THEN 1 ELSE 0 END) AS unchanged_cnt,
             SUM(CASE WHEN CAST(cash_quota_amount_now_after AS DOUBLE) < CAST(jq_credit_quota AS DOUBLE) THEN 1 ELSE 0 END) AS down_cnt
      FROM {TABLE}
      WHERE jq_date='2026-08-12'
        AND CAST(jq_credit_quota AS DOUBLE) >= 0
        AND CAST(cash_quota_amount_now_after AS DOUBLE) >= 0
      GROUP BY str_type, app_flow_flag, te_flag
      ORDER BY str_type, app_flow_flag, te_flag
    """
    group_sql = f"""
      SELECT jq_customer_flag, defq_use_rate_level, separate_installment_loan_cnt,
             jq_mix_score_bin, cash_remark1, te_flag,
             COUNT(1) AS cnt,
             SUM(CAST(jq_credit_quota AS DOUBLE)) AS before_sum_cent,
             SUM(CAST(cash_quota_amount_now_after AS DOUBLE)) AS after_sum_cent,
             SUM(CAST(cash_quota_amount_now_after AS DOUBLE) - CAST(jq_credit_quota AS DOUBLE)) AS increment_cent
      FROM {TABLE}
      WHERE jq_date='2026-08-12' AND str_type='new'
        AND CAST(jq_credit_quota AS DOUBLE) >= 0
        AND CAST(cash_quota_amount_now_after AS DOUBLE) >= 0
      GROUP BY jq_customer_flag, defq_use_rate_level, separate_installment_loan_cnt,
               jq_mix_score_bin, cash_remark1, te_flag
    """
    strategy_summary = query(odps, summary_sql)
    groups = query(odps, group_sql)

    policy = policy_lookup()
    plan = plan_rates()
    total = empty_metric()
    by_te: dict[str, dict[str, float | int]] = defaultdict(empty_metric)
    by_score: dict[str, dict[str, float | int]] = defaultdict(empty_metric)
    direction: dict[str, dict[str, float | int]] = defaultdict(empty_metric)

    for row in groups:
        count = int(row["cnt"])
        before = float(row["before_sum_cent"])
        after = float(row["after_sum_cent"])
        increment = float(row["increment_cent"])
        actual_rate = as_float(row["cash_remark1"])
        key = lookup_key(
            row["jq_customer_flag"], row["defq_use_rate_level"],
            row["separate_installment_loan_cnt"], row["jq_mix_score_bin"],
        )
        expected_rate = policy.get(key) if key else None
        planned = plan.get(key) if key else None
        plan_a_rate = float(planned["plan_a_rate"]) if planned else 0.0
        plan_b_rate = float(planned["plan_b_rate"]) if planned else 0.0
        is_plan_abc = bool(planned and planned["plan_abc_eligible"])
        score = normalized_code(row["jq_mix_score_bin"])

        targets = [total, by_te[str(row["te_flag"])], by_score[score]]
        if expected_rate is None:
            direction_name = "no_lookup"
        elif actual_rate is None:
            direction_name = "missing_actual_rate"
        elif abs(actual_rate - expected_rate) < 1e-10:
            direction_name = "match"
        elif actual_rate < expected_rate:
            direction_name = "actual_rate_lower_than_table"
        else:
            direction_name = "actual_rate_higher_than_table"
        targets.append(direction[direction_name])

        for metric in targets:
            metric["cnt"] = int(metric["cnt"]) + count
            metric["before_sum_cent"] = float(metric["before_sum_cent"]) + before
            metric["after_sum_cent"] = float(metric["after_sum_cent"]) + after
            metric["actual_increment_cent"] = float(metric["actual_increment_cent"]) + increment
            if actual_rate is not None:
                metric["actual_rate_unconstrained_increment_cent"] = float(metric["actual_rate_unconstrained_increment_cent"]) + before * actual_rate
            if expected_rate is None:
                metric["no_lookup_cnt"] = int(metric["no_lookup_cnt"]) + count
            else:
                metric["coefficient_table_increment_cent"] = float(metric["coefficient_table_increment_cent"]) + before * expected_rate
                if actual_rate is not None:
                    rate_diff = expected_rate - actual_rate
                    metric["rate_table_shortfall_cent"] = float(metric["rate_table_shortfall_cent"]) + before * rate_diff
                    if abs(rate_diff) < 1e-10:
                        metric["coefficient_match_cnt"] = int(metric["coefficient_match_cnt"]) + count
                    else:
                        metric["coefficient_mismatch_cnt"] = int(metric["coefficient_mismatch_cnt"]) + count
                        if actual_rate < expected_rate:
                            metric["actual_rate_below_table_cnt"] = int(metric["actual_rate_below_table_cnt"]) + count
                        else:
                            metric["actual_rate_above_table_cnt"] = int(metric["actual_rate_above_table_cnt"]) + count
            if is_plan_abc:
                metric["plan_a_increment_cent"] = float(metric["plan_a_increment_cent"]) + before * plan_a_rate
                metric["plan_b_increment_cent"] = float(metric["plan_b_increment_cent"]) + before * plan_b_rate
                if expected_rate is not None:
                    metric["plan_a_vs_table_gap_cent"] = float(metric["plan_a_vs_table_gap_cent"]) + before * (plan_a_rate - expected_rate)
            if str(row["te_flag"]) == "0" and actual_rate is not None and actual_rate > 0:
                metric["positive_config_no_raise_cnt"] = int(metric["positive_config_no_raise_cnt"]) + count

    valid_summary_total = empty_metric()
    for row in strategy_summary:
        count = int(row["cnt"])
        valid_summary_total["cnt"] = int(valid_summary_total["cnt"]) + count
        valid_summary_total["before_sum_cent"] = float(valid_summary_total["before_sum_cent"]) + float(row["before_sum_cent"])
        valid_summary_total["after_sum_cent"] = float(valid_summary_total["after_sum_cent"]) + float(row["after_sum_cent"])
        valid_summary_total["actual_increment_cent"] = float(valid_summary_total["actual_increment_cent"]) + float(row["increment_cent"])

    output = {
        "scope": {
            "business_date": "2026-08-12",
            "new_strategy_scope": "str_type=new",
            "target_average_limit_yuan": 50000,
            "amount_unit": "cent",
        },
        "full_valid_population": enrich(valid_summary_total),
        "strategy_execution_summary": strategy_summary,
        "new_strategy_total": enrich(total),
        "new_strategy_by_te_flag": {key: enrich(value) for key, value in by_te.items()},
        "new_strategy_by_score_bin": {key: enrich(value) for key, value in by_score.items()},
        "new_strategy_by_coefficient_direction": {key: enrich(value) for key, value in direction.items()},
    }
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
