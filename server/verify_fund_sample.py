"""用资金归结样本 Excel 核验计算结果。

用法:
    python server/verify_fund_sample.py --result result.json --expected sample.xlsx

result.json 可以是 runner 的 payload，也可以是 ``{"result": payload}``。
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


def _rows(sheet: Any) -> list[list[Any]]:
    return [list(row) for row in sheet.iter_rows(values_only=True) if any(value is not None for value in row)]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or _text(value) == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _same(left: Any, right: Any, tolerance: float = 1e-6) -> bool:
    left_number, right_number = _number(left), _number(right)
    if left_number is not None and right_number is not None:
        return math.isclose(left_number, right_number, rel_tol=1e-7, abs_tol=tolerance)
    return _text(left) == _text(right)


def _check_numeric(
    failed: list[str], sheet_name: str, label: str, expected: Any, actual: Any, tolerance: float = 1e-6
) -> None:
    if not _same(expected, actual, tolerance):
        failed.append(f"{sheet_name}/{label}: Excel={expected!r}, result={actual!r}")


def _find_header(rows: list[list[Any]], name: str) -> tuple[int, list[Any]] | None:
    for index, row in enumerate(rows):
        if name in row:
            return index, row
    return None


def _cell_by_header(row: list[Any], header: list[Any], name: str) -> Any:
    try:
        return row[header.index(name)]
    except (ValueError, IndexError):
        return None


def _verify_summary(payload: dict[str, Any], book: Any, failed: list[str], checked: list[str]) -> None:
    for sheet_name in ("样本校验", "结论摘要"):
        if sheet_name not in book.sheetnames:
            continue
        rows = _rows(book[sheet_name])
        for row in rows:
            if not row:
                continue
            label = _text(row[0])
            value = row[1] if len(row) > 1 else None
            if label == "有效日期数":
                _check_numeric(failed, sheet_name, label, payload["meta"]["date_count"], value)
                checked.append(f"{sheet_name}/{label}")
            elif label == "日期范围":
                expected = f"{payload['meta']['date_start']} ~ {payload['meta']['date_end']}"
                _check_numeric(failed, sheet_name, label, expected, value)
                checked.append(f"{sheet_name}/{label}")
            elif label in ("原始行数", "总订单数"):
                _check_numeric(failed, sheet_name, label, payload["summary"]["total_order_count"], value)
                checked.append(f"{sheet_name}/{label}")
            elif label == "异常订单数":
                _check_numeric(failed, sheet_name, label, payload["summary"]["abnormal_order_count"], value)
                checked.append(f"{sheet_name}/{label}")


def _verify_windows(payload: dict[str, Any], book: Any, failed: list[str], checked: list[str]) -> None:
    if "窗口总览" not in book.sheetnames:
        failed.append("缺少工作表：窗口总览")
        return
    rows = _rows(book["窗口总览"])
    header_info = _find_header(rows, "窗口")
    if header_info is None:
        failed.append("窗口总览：缺少表头")
        return
    header_index, header = header_info
    excel_by_label = {_text(row[0]): row for row in rows[header_index + 1:] if row and _text(row[0])}
    fields = [
        ("观察天数", "observation_days"),
        ("基准天数", "baseline_days"),
        ("基准异常订单", "baseline_abnormal_count"),
        ("观察异常订单", "observation_abnormal_count"),
        ("异常订单量增长倍数", "growth_factor"),
        ("基准总订单", "baseline_total_count"),
        ("观察总订单", "observation_total_count"),
        ("基准异常率", "baseline_rate"),
        ("观察异常率", "observation_rate"),
        ("异常率提升倍数", "rate_lift_factor"),
        ("异常率变化_bp", "rate_change_bp"),
        ("z-score", "z_score"),
    ]
    for metric in payload.get("window_overview", []):
        label = _text(metric.get("label"))
        row = excel_by_label.get(label)
        if row is None:
            failed.append(f"窗口总览/{label}: Excel 缺少该窗口")
            continue
        for excel_name, payload_name in fields:
            _check_numeric(failed, "窗口总览", f"{label}/{excel_name}", _cell_by_header(row, header, excel_name), metric.get(payload_name), tolerance=1e-5)
        checked.append(f"窗口总览/{label}")


def _verify_daily(payload: dict[str, Any], book: Any, failed: list[str], checked: list[str]) -> None:
    if "大盘日趋势" not in book.sheetnames:
        failed.append("缺少工作表：大盘日趋势")
        return
    rows = _rows(book["大盘日趋势"])
    header_info = _find_header(rows, "日期")
    if header_info is None:
        failed.append("大盘日趋势：缺少表头")
        return
    header_index, header = header_info
    excel_by_date = {_text(row[0])[:10]: row for row in rows[header_index + 1:] if row and _text(row[0])}
    for point in payload.get("daily_trend", []):
        date = _text(point.get("date"))[:10]
        row = excel_by_date.get(date)
        if row is None:
            failed.append(f"大盘日趋势/{date}: Excel 缺少该日期")
            continue
        _check_numeric(failed, "大盘日趋势", f"{date}/异常订单数", _cell_by_header(row, header, "异常订单数"), point.get("abnormal_order_count"), tolerance=1e-5)
        _check_numeric(failed, "大盘日趋势", f"{date}/总订单数", _cell_by_header(row, header, "总订单数"), point.get("order_count"), tolerance=1e-5)
        _check_numeric(failed, "大盘日趋势", f"{date}/异常率", _cell_by_header(row, header, "异常率"), point.get("abnormal_rate"), tolerance=1e-8)
    if len(excel_by_date) != len(payload.get("daily_trend", [])):
        failed.append(f"大盘日趋势/日期数: Excel={len(excel_by_date)}, result={len(payload.get('daily_trend', []))}")
    checked.append("大盘日趋势")


def _verify_expert(payload: dict[str, Any], book: Any, failed: list[str], checked: list[str]) -> None:
    sheet_name = "专家_强制单维结果"
    if sheet_name not in book.sheetnames:
        failed.append(f"缺少工作表：{sheet_name}")
        return
    rows = _rows(book[sheet_name])
    header_info = _find_header(rows, "维度路径")
    if header_info is None:
        failed.append(f"{sheet_name}：缺少表头")
        return
    header_index, header = header_info
    excel_paths = {_text(row[header.index("维度路径")]) for row in rows[header_index + 1:] if len(row) > header.index("维度路径") and _text(row[header.index("维度路径")])}
    result_paths = {str(record.get("path")) for record in payload.get("expert", {}).get("single", [])}
    if excel_paths != result_paths:
        failed.append(f"{sheet_name}/维度路径: Excel={sorted(excel_paths)!r}, result={sorted(result_paths)!r}")
    checked.append(f"{sheet_name}/维度路径")


def _verify_alert_state(payload: dict[str, Any], book: Any, failed: list[str], checked: list[str]) -> None:
    sheet_name = "合并预警结果"
    if sheet_name not in book.sheetnames:
        failed.append(f"缺少工作表：{sheet_name}")
        return
    rows = _rows(book[sheet_name])
    no_alert = any(row and "本轮无正式预警" in _text(row[0]) for row in rows)
    if no_alert and payload.get("merged_alerts"):
        failed.append(f"{sheet_name}: Excel 标记无正式预警，result={len(payload['merged_alerts'])} 条")
    if not no_alert and not payload.get("merged_alerts"):
        failed.append(f"{sheet_name}: Excel 存在正式预警内容，result=0 条")
    checked.append(sheet_name)


def verify_payload_against_workbook(payload: dict[str, Any], workbook_path: Path | str) -> dict[str, Any]:
    """返回可供 CI/人工查看的核验报告，不修改 Excel。"""
    workbook = load_workbook(Path(workbook_path), data_only=True, read_only=True)
    failed: list[str] = []
    checked: list[str] = []
    _verify_summary(payload, workbook, failed, checked)
    _verify_windows(payload, workbook, failed, checked)
    _verify_daily(payload, workbook, failed, checked)
    _verify_expert(payload, workbook, failed, checked)
    _verify_alert_state(payload, workbook, failed, checked)
    return {"ok": not failed, "checked": checked, "failed": failed}


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="核验资金归结结果 JSON 与 Excel 模板是否一致")
    parser.add_argument("--result", required=True, type=Path, help="runner 输出 JSON")
    parser.add_argument("--expected", required=True, type=Path, help="Excel 模板运行结果")
    args = parser.parse_args(list(argv) if argv is not None else None)
    raw = json.loads(args.result.read_text(encoding="utf-8"))
    payload = raw.get("result", raw)
    report = verify_payload_against_workbook(payload, args.expected)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
