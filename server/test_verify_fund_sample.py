"""资金样本 Excel 核验器测试。"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook

SERVER_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SERVER_DIR))

from test_fund_attribution import OfflineRunner, _detail_frame  # noqa: E402
from verify_fund_sample import verify_payload_against_workbook  # noqa: E402


def _payload() -> dict:
    config = json.loads((SERVER_DIR / "config" / "fund_monitor_config.json").read_text(encoding="utf-8"))
    config["dimensions"] = ["traffic_channel", "loan_cust_lifecycle_stage"]
    config["field_labels"] = {"traffic_channel": "流量渠道", "loan_cust_lifecycle_stage": "客户生命周期"}
    return OfflineRunner(config, _detail_frame(days=20, spike=True), partition="20260903").run()


def _write_reference(path: Path, payload: dict) -> None:
    book = Workbook()
    book.remove(book.active)
    windows = book.create_sheet("窗口总览")
    windows.append(["窗口", "观察期", "基准期", "观察天数", "基准天数", "基准异常订单", "观察异常订单", "异常订单量增长倍数", "基准总订单", "观察总订单", "基准异常率", "观察异常率", "异常率提升倍数", "异常率变化_bp", "z-score"])
    for item in payload["window_overview"]:
        windows.append([item["label"], item["observation_start"], item["baseline_start"], item["observation_days"], item["baseline_days"], item["baseline_abnormal_count"], item["observation_abnormal_count"], item["growth_factor"], item["baseline_total_count"], item["observation_total_count"], item["baseline_rate"], item["observation_rate"], item["rate_lift_factor"], item["rate_change_bp"], item["z_score"]])
    daily = book.create_sheet("大盘日趋势")
    daily.append(["日期", "异常订单数", "总订单数", "异常率"])
    for item in payload["daily_trend"]:
        daily.append([item["date"], item["abnormal_order_count"], item["order_count"], item["abnormal_rate"]])
    merged = book.create_sheet("合并预警结果")
    merged.append(["本轮无正式预警"] if not payload["merged_alerts"] else ["路径"])
    expert = book.create_sheet("专家_强制单维结果")
    expert.append(["排名", "归因来源", "维度层级", "维度路径", "最终预警等级"])
    for index, item in enumerate(payload["expert"]["single"], start=1):
        expert.append([index, item["source"], item["layer"], item["path"], item["level_label"]])
    book.save(path)


def test_verifier_accepts_matching_workbook(tmp_path: Path):
    payload = _payload()
    workbook = tmp_path / "reference.xlsx"
    _write_reference(workbook, payload)
    report = verify_payload_against_workbook(payload, workbook)
    assert report["ok"] is True
    assert report["failed"] == []


def test_verifier_reports_metric_mismatch(tmp_path: Path):
    payload = _payload()
    workbook = tmp_path / "reference.xlsx"
    _write_reference(workbook, payload)
    payload["window_overview"][0]["observation_abnormal_count"] += 1
    report = verify_payload_against_workbook(payload, workbook)
    assert report["ok"] is False
    assert any("近1天" in item for item in report["failed"])
